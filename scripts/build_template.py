"""풀모델 정본 템플릿(`valstudio-full-v1`) 결정론 빌드.

`ValStudio_DCF_Template.xlsx`(26시트)를 **원본 서식 + 코드**에서 재생성한다. 워크북 자체는
`.gitignore`(`*.xlsx`) 대상이라 레포에 없으므로, **파생 절차를 코드로 남기는 것**이 이 스크립트의
목적이다 — 화면이 바뀌거나 원본이 갱신돼도 같은 결과를 다시 만들 수 있고, 셀 좌표 계약
(`backend/excel/fullmodel_layout.py`)과 템플릿이 **같은 소스에서 나와** 어긋나지 않는다.

파생 절차(원본 → 정본):
  ① 외부 브랜딩 문자열·임베드 이미지 제거          ② 죽은 시트 제거(피참조 0 실측)
  ③ 유사회사 퍼널 시트 이관(수식·서식 보존)         ④ Comps=Trading Comps 수식째 이관 후 데이터 blank
  ⑤ Research(조사 백데이터 10섹션)·rFS 신설         ⑥ 거시 RAW 분리(r* 3종)
  ⑦ **상향 배선**(세부→요약→DCF 스파인, WACC 빌드업) ⑧ Format·_DART_MAP·_VS_STATE

사용:
  python scripts/build_template.py                 # 빌드 + 검증
  python scripts/build_template.py --out other.xlsx
  python scripts/build_template.py --verify-only   # 기존 산출물만 재검증

원본 3종이 없으면 **명시 에러**로 중단한다(조용히 반쪽 템플릿을 만들지 않는다).
"""
from __future__ import annotations

import re
import sys
from copy import copy
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
except ImportError:                                    # 빌드 전용 의존(엔진은 stdlib)
    raise SystemExit("openpyxl 필요: pip install openpyxl")

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "backend"))

# ── 원본 3종 ────────────────────────────────────────────────────────────────
SRC_BASE = ROOT / "(DCF)성명_평가대상회사명_DCF Model.xlsx"     # 스파인·EBIT·FA·WC·WACC 골격
SRC_PEER = ROOT / "할인율 서식.xlsx"                            # 유사회사 4-step 퍼널
SRC_COMPS = Path(r"D:\Valuation\모델러스엑셀\5.4(COMPLETED).xlsx")  # Trading Comps·RAW 패턴
OUT_DEFAULT = ROOT / "ValStudio_DCF_Template.xlsx"

LAYOUT_ID = "valstudio-full-v1"
TITLE = "Val-Studio | DCF Valuation Template v1"

# ── 디자인 토큰(독자 서식) ──────────────────────────────────────────────────
HEAD, BAND, GREY = "FF1E293B", "FFF1F5F9", "FF64748B"
BLUE, BLACK, GREEN, YEL = "FF1D4ED8", "FF0F172A", "FF15803D", "FFFEF08A"
THIN = Side(style="thin", color="FFCBD5E1")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# ── 좌표 상수 ───────────────────────────────────────────────────────────────
FCOLS = list("MNOPQ")            # DCF 추정기간 5개년 (H:L=실적, R=Terminal)
ACOLS = list("EFGHIJ")           # Assumption 거시 블록 6개년
RCOLS = list("CDEFGH")           # r* 시트 연도열
YEARS = list(range(2023, 2029))

BRAND = re.compile(r"MSVALUE[^\"<]*|기업가치평가\s*연수\s*\d*기", re.I)
DEAD_SHEETS = ["유사회사FS", "BackData"]      # 실측: 내용 없음·피참조 0(강의 안내 문구뿐)
PEER_RENAME = {
    "유사기업선정>>": "PEER >>", "선정과정": "Peer_선정과정",
    "Step0. 평가대상회사 리서치": "Peer_S0_대상리서치",
    "Step1. 모집단 선정": "Peer_S1_모집단",
    "Step2. 사업유사성 검토": "Peer_S2_사업유사성",
    "Step3. 매출비중검토": "Peer_S3_매출비중",
    "할인율>>": None,            # 구분자 중복 → 제외
    "WACC": "WACC_Ref",          # 베이스와 충돌 → 중복도 판정 후 결정
}
MACRO = {                         # r* 시트 → (Assumption 행, 라벨, 설명, 출처)
    "rGDP": (9, "실질 GDP 성장률", "Real GDP growth (%)", "한국은행 ECOS / EIU"),
    "rInflation": (10, "소비자물가 상승률", "Consumer Price Inflation (av, %)", "한국은행 ECOS / EIU"),
    "rWage": (11, "명목임금 상승률", "Average nominal wages (%)", "고용노동부 / EIU"),
}
TAB_ORDER = [
    "Assumption", "DCF", "EBIT", "FA", "WC", "WACC", "H_FS", "상각비계산", "Comps",
    "PEER >>", "Peer_선정과정", "Peer_S0_대상리서치", "Peer_S1_모집단",
    "Peer_S2_사업유사성", "Peer_S3_매출비중",
    "RESEARCH >>", "Research",
    "RAW >>", "rFS", "rGDP", "rInflation", "rWage", "rTrading",
    "Format", "_DART_MAP", "_VS_STATE",
]

log: list[str] = []


def say(msg: str) -> None:
    log.append(msg)
    print(f"  {msg}")


# ═══════════════════════════════════════════════════════════════════════════
# 공통 헬퍼
# ═══════════════════════════════════════════════════════════════════════════
class Pen:
    """행 커서를 들고 섹션·표를 찍는다 — 행 번호 수작업 계산을 없애 오프바이원을 차단."""

    def __init__(self, ws, r: int = 1):
        self.ws, self.r = ws, r

    def section(self, title: str) -> "Pen":
        self.ws.cell(row=self.r, column=2, value="x").font = Font(color="FFCBD5E1", size=8)
        c = self.ws.cell(row=self.r, column=3, value=title)
        c.font = Font(bold=True, size=12, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD)
        self.r += 2
        return self

    def block(self, label: str) -> "Pen":
        self.ws.cell(row=self.r, column=3, value=label).font = Font(bold=True, size=10, color=HEAD)
        self.r += 1
        return self

    def table(self, headers: list[str], rows: int = 6, col: int = 3) -> "Pen":
        for j, h in enumerate(headers):
            c = self.ws.cell(row=self.r, column=col + j, value=h)
            c.font = Font(bold=True, size=9, color="FFFFFFFF")
            c.fill = PatternFill("solid", fgColor="FF475569")
            c.alignment = Alignment(horizontal="center")
            c.border = BOX
        self.r += 1
        for _ in range(rows):
            for j in range(len(headers)):
                cc = self.ws.cell(row=self.r, column=col + j)
                cc.font = Font(color=BLUE)
                cc.border = BOX
            self.r += 1
        return self

    def kv(self, labels: list[str], col: int = 3) -> "Pen":
        for lab in labels:
            self.ws.cell(row=self.r, column=col, value=lab).border = BOX
            for j in (1, 2):
                cc = self.ws.cell(row=self.r, column=col + j)
                cc.font = Font(color=BLUE)
                cc.border = BOX
            self.r += 1
        return self

    def note(self, text: str, fill: str | None = None) -> "Pen":
        c = self.ws.cell(row=self.r, column=3, value=text)
        c.font = Font(size=9, color=GREY)
        if fill:
            c.fill = PatternFill("solid", fgColor=fill)
            c.font = Font(size=9, color=BLACK)
        self.r += 1
        return self

    def src(self, hint: str = "[출처 URL·조회일자]") -> "Pen":
        c = self.ws.cell(row=self.r, column=3, value=hint)
        c.font = Font(color=BLUE, size=9, underline="single")
        self.r += 2
        return self

    def space(self, n: int, label: str) -> "Pen":
        c = self.ws.cell(row=self.r, column=3, value=label)
        c.font = Font(size=9, color=GREY, italic=True)
        c.fill = PatternFill("solid", fgColor="FFF8FAFC")
        self.r += n
        return self

    def gap(self, n: int = 1) -> "Pen":
        self.r += n
        return self


def hdr(ws, row: int, col: int, labels: list[str], width: int | None = None) -> None:
    for j, h in enumerate(labels):
        c = ws.cell(row=row, column=col + j, value=h)
        c.font = Font(bold=True, size=9, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BOX
        if width:
            ws.column_dimensions[openpyxl.utils.get_column_letter(col + j)].width = width


def title_of(ws, text: str, sub: str = "") -> None:
    ws["B2"] = text
    ws["B2"].font = Font(bold=True, size=13, color=HEAD)
    if sub:
        ws["B3"] = sub
        ws["B3"].font = Font(size=9, color=GREY)


def copy_sheet(src, dst) -> int:
    """셀 값·수식·서식·병합·열너비 이관(openpyxl 은 워크북 간 시트 복사 API 가 없다)."""
    n = 0
    for row in src.iter_rows():
        for c in row:
            if c.value is None and not c.has_style:
                continue
            t = dst.cell(row=c.row, column=c.column, value=c.value)
            if c.has_style:
                t.font, t.fill = copy(c.font), copy(c.fill)
                t.border, t.alignment = copy(c.border), copy(c.alignment)
                t.number_format, t.protection = c.number_format, copy(c.protection)
            n += 1
    for rng in src.merged_cells.ranges:
        dst.merge_cells(str(rng))
    for k, d in src.column_dimensions.items():
        dst.column_dimensions[k].width, dst.column_dimensions[k].hidden = d.width, d.hidden
    for k, d in src.row_dimensions.items():
        dst.row_dimensions[k].height = d.height
    dst.sheet_view.showGridLines = src.sheet_view.showGridLines
    return n


def blank_values(ws, ranges: list[str]) -> int:
    """수식은 남기고 하드값(원본 회사 데이터)만 제거 — 구조·계산은 보존."""
    n = 0
    for ref in ranges:
        got = ws[ref]
        rows = got if isinstance(got, tuple) else ((got,),)     # 단일 셀은 Cell 반환
        for row in rows:
            row = row if isinstance(row, tuple) else (row,)
            for c in row:
                if c.value is not None and not (isinstance(c.value, str)
                                                and c.value.startswith("=")):
                    c.value = None
                    n += 1
    return n


def put(ws, ref: str, formula: str, color: str = BLACK, bold: bool = False) -> bool:
    """**빈 셀에만** 기록 — 원본 수식/값은 절대 덮어쓰지 않는다(파생의 안전장치)."""
    c = ws[ref]
    if c.value is not None:
        return False
    c.value = formula
    c.font = Font(color=color, bold=bold)
    return True


# ═══════════════════════════════════════════════════════════════════════════
# 페이즈 ① 정제 — 브랜딩·이미지·죽은 시트
# ═══════════════════════════════════════════════════════════════════════════
def phase_clean(wb) -> None:
    n = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and BRAND.search(c.value):
                    c.value = BRAND.sub("", c.value).strip() or TITLE
                    c.font = Font(bold=True, size=10, color=HEAD)
                    n += 1
    say(f"브랜딩 문자열 치환 {n}건")

    dropped = 0
    for ws in wb.worksheets:                       # 제3자 데이터 스크린샷(단말기 캡처) 제거
        if getattr(ws, "_images", []):
            dropped += len(ws._images)
            ws._images = []
    say(f"임베드 이미지 제거 {dropped}건")

    for name in DEAD_SHEETS:
        if name in wb.sheetnames:
            del wb[name]
    say(f"죽은 시트 제거: {', '.join(DEAD_SHEETS)}")


# ═══════════════════════════════════════════════════════════════════════════
# 페이즈 ② 퍼널 이관 — 할인율 서식
# ═══════════════════════════════════════════════════════════════════════════
def _fingerprint(ws, limit: int = 400) -> list[str]:
    out = []
    for row in ws.iter_rows(min_row=1, max_row=60):
        for c in row:
            if isinstance(c.value, str) and c.value.strip():
                out.append(c.value.strip()[:30])
            if len(out) >= limit:
                return out
    return out


def phase_peer(wb) -> None:
    src = openpyxl.load_workbook(SRC_PEER, data_only=False)

    # 두 워크북 모두 WACC 시트를 갖는다 — 라벨 지문으로 중복 판정(실측 Jaccard 73%).
    fa, fb = _fingerprint(wb["WACC"]), _fingerprint(src["WACC"])
    jac = len(set(fa) & set(fb)) / max(1, len(set(fa) | set(fb)))
    skip_wacc = jac > 0.7
    say(f"WACC 중복도 {jac:.0%} → {'할인율서식 WACC 제외' if skip_wacc else 'WACC_Ref 편입'}")

    for name in src.sheetnames:
        new = PEER_RENAME.get(name, name)
        if new is None or (name == "WACC" and skip_wacc):
            continue
        ws = wb.create_sheet(new)
        n = copy_sheet(src[name], ws)
        ws.sheet_properties.tabColor = "1D4ED8" if new.startswith("Peer") else "94A3B8"
        say(f"이관 {name} → {new} ({n:,}셀)")


# ═══════════════════════════════════════════════════════════════════════════
# 페이즈 ③ Comps — Trading Comps 수식째 이관
# ═══════════════════════════════════════════════════════════════════════════
def phase_comps(wb) -> None:
    """손으로 다시 그리지 않는다 — LTM/FTM 2축·NM 처리·High/Low/Median/Average 가
    이미 검증된 형태이므로 수식째 복사하고 **회사 고유 데이터만** 비운다."""
    ref = openpyxl.load_workbook(SRC_COMPS, data_only=False)

    comps = wb.create_sheet("Comps")
    n1 = copy_sheet(ref["Trading"], comps)
    comps.sheet_properties.tabColor = "1D4ED8"
    b1 = blank_values(comps, ["D3", "N14:AP18", "Q6:Q7", "L5:L7"])
    comps["D3"] = "[평가대상회사]"
    comps["D3"].font = Font(bold=True, size=12, color=HEAD)

    rt = wb.create_sheet("rTrading")
    n2 = copy_sheet(ref["rTrading"], rt)
    rt.sheet_properties.tabColor = "CBD5E1"
    b2 = blank_values(rt, ["A3:Y13"])
    say(f"Comps ← Trading {n1:,}셀(하드값 {b1} 제거) / rTrading {n2:,}셀({b2} 제거)")


# ═══════════════════════════════════════════════════════════════════════════
# 페이즈 ④ 신설 시트 — Research · rFS · r* · Format · 구분자
# ═══════════════════════════════════════════════════════════════════════════
def build_research(wb):
    """조사 백데이터 10섹션. 실무 기업리서치 양식 실측 구조(섹션 마커 + 서사·숫자·출처 동거)."""
    ws = wb.create_sheet("Research")
    ws.sheet_properties.tabColor = "1D4ED8"
    for col, w in (("B", 3), ("C", 30), ("D", 40), ("E", 14), ("F", 14),
                   ("G", 14), ("H", 26), ("I", 22), ("J", 14), ("K", 12)):
        ws.column_dimensions[col].width = w
    ws["C1"] = "[평가대상회사] 리서치"
    ws["C1"].font = Font(bold=True, size=14, color=HEAD)
    ws["C2"] = ("조사 백데이터 — 서사·숫자·출처를 한 곳에. 하류 시트(EBIT·Assumption)는 "
                "여기 숫자를 참조하고, 판단 근거는 여기 서사를 읽는다.")
    ws["C2"].font = Font(size=9, color=GREY)

    p = Pen(ws, 4)
    p.section("1. Summary — 투자포인트")
    p.note("한 줄씩. 각 포인트는 아래 섹션 중 하나가 뒷받침해야 한다(무근거 포인트 금지).")
    for _ in range(5):
        ws.cell(row=p.r, column=3).font = Font(color=BLUE)
        p.r += 1
    p.src().gap()

    p.section("2. 회사 개요")
    p.block("[회사개요]").table(["구분", "내용"], rows=0)
    p.kv(["설립일", "창업자", "대표이사", "사업개요", "본사 소재지", "홈페이지",
          "상장시장·상장일", "결산월", "신용등급", "감사인", "감사의견"])
    p.gap().block("[주요 종속회사]").table(["회사", "지분율", "사업 개요"], rows=5)
    p.gap().block("[주주 구성]").table(["주주", "지분율", "비고"], rows=6)
    p.kv(["유통주식 비율", "발행주식수", "자기주식"])
    p.note("⚠ 발행주식수 ≠ 유통주식수 — DCF 주당가치 분모를 어느 쪽으로 쓸지 여기서 확정.", YEL)
    p.src("[DART 사업보고서 URL] / [FnGuide URL]").gap()

    p.section("3. 지배구조 · 사업 구조도")
    p.space(14, "▸ 지분율/사업 구조도 이미지 붙여넣기 영역").src().gap()

    p.section("4. 제품·서비스별 매출 및 비중")
    p.note("→ EBIT 시트 '매출 세부 추정내역'(사업1~4)의 직접 근거. 여기 부문 구분이 곧 추정 단위.")
    p.block("[사업부문별 매출]").table(
        ["부문명", "직전연도 매출", "비중", "전기 매출", "전기 비중", "수량(Q) 단위", "단가(P) 단위"],
        rows=6)
    first = p.r - 6
    ws.cell(row=p.r, column=3, value="합 계").font = Font(bold=True)
    ws.cell(row=p.r, column=4, value=f"=SUM(D{first}:D{p.r-1})").font = Font(bold=True)
    ws.cell(row=p.r, column=5, value=f"=SUM(E{first}:E{p.r-1})").font = Font(bold=True)
    p.r += 2
    p.block("[종속회사별 제품 현황]").table(["회사", "사업 부문", "매출액", "비율"], rows=8)
    p.src().gap()

    p.section("5. 주요 제품 설명")
    p.block("[제품별]").table(["제품명", "내용"], rows=0)
    p.kv(["제품1 설명", "  매출 출처(전방)", "  시장 점유율", "제품2 설명",
          "  매출 출처(전방)", "  시장 점유율", "제품3 설명", "  매출 출처(전방)", "  시장 점유율"])
    p.gap().space(10, "▸ 제품 사진·도해 붙여넣기 영역").src().gap()

    p.section("6. Value Chain · 산업 구조")
    p.block("[밸류체인]").space(10, "▸ 밸류체인 도해 붙여넣기 영역")
    p.block("[공정·단계별 설명]").table(["구분", "내용"], rows=0)
    p.kv(["전방산업", "후방산업", "당사 위치", "진입장벽", "교체비용"]).src().gap()

    p.section("7. 경쟁 구도")
    p.block("[경쟁사 및 점유율]").table(["경쟁사", "국가", "점유율", "강점", "약점", "비고"], rows=7)
    p.kv(["경쟁 강도 판단", "당사 차별점"]).src().gap()

    p.section("8. 시장 규모 · 성장률")
    p.block("[시장 데이터]").table(["구분", "직전", "전망1", "전망2", "전망3", "CAGR", "출처"], rows=6)
    p.note("→ Assumption 시트 '성장률 전망'의 근거. 거시 지표는 rGDP·rInflation·rWage 참조.")
    p.src().gap()

    p.section("9. 고객·매출처")
    p.block("[주요 고객]").table(["고객사", "매출 비중", "거래 기간", "계약 형태", "비고"], rows=6)
    p.kv(["매출처 집중도 판단"]).src().gap()

    p.section("10. 자료 출처 로그")
    p.note("여기 없는 자료는 가정 근거로 쓰지 않는다(provenance 게이트).")
    p.block("[출처 목록]").table(
        ["자료명", "발행처", "발행일", "조회일자", "URL/경로", "활용 단계"], rows=14)
    return ws


def build_rfs(wb):
    ws = wb.create_sheet("rFS")
    ws.sheet_properties.tabColor = "CBD5E1"
    title_of(ws, "재무제표 원문 (RAW)",
             "DART 원문을 가공 없이 붙여넣는 곳. 정규화·재분류는 H_FS 가 담당한다 "
             "— 여기서 숫자를 고치면 원문 대조가 불가능해진다.")
    ws["B5"] = "[출처 URL·보고서명·조회일자]"
    ws["B5"].font = Font(color=BLUE, size=9, underline="single")
    for col, label in (("B", "재무상태표"), ("H", "손익계산서"), ("N", "현금흐름표")):
        c = ws[f"{col}7"]
        c.value = f"[{label}]"
        c.font = Font(bold=True, size=11, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD)
        ws.column_dimensions[col].width = 34
    ws["B8"] = "▸ 계정과목 | 연도별 금액을 그대로 붙여넣기 (단위 표기 필수)"
    ws["B8"].font = Font(size=9, color=GREY)
    return ws


def build_macro_raw(wb) -> None:
    """거시 RAW 분리 — Assumption 하드값을 여기로 **이사**시키고 Assumption 은 참조로 전환."""
    asm = wb["Assumption"]
    for name, (arow, label, desc, source) in MACRO.items():
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = "CBD5E1"
        ws["B2"] = label
        ws["B2"].font = Font(bold=True, size=12, color=HEAD)
        ws["B3"] = desc
        ws["B4"] = f"출처: {source}"
        ws["B5"] = "출처·조회일자 — Assumption 원 하드값 이사분. 갱신 시 여기만 수정"
        ws["B5"].font = Font(color=BLUE, size=9)
        ws["B7"], ws["B8"], ws["B9"] = "연도", "값(%)", "구분(A=실적/F=전망)"
        ws["B7"].font = Font(bold=True)
        ws.column_dimensions["B"].width = 24
        for j, (rc, y) in enumerate(zip(RCOLS, YEARS)):
            ws[f"{rc}7"] = y
            ws[f"{rc}7"].font = Font(bold=True)
            ws[f"{rc}7"].fill = PatternFill("solid", fgColor=BAND)
            ws[f"{rc}8"] = asm[f"{ACOLS[j]}{arow}"].value        # 값 이사(수정 아님)
            ws[f"{rc}8"].font = Font(color=BLUE)
            ws[f"{rc}9"] = "A" if y == YEARS[0] else "F"
            ws[f"{rc}9"].font = Font(color=BLUE)
        for j, ac in enumerate(ACOLS):                            # Assumption → 초록 참조
            asm[f"{ac}{arow}"] = f"={name}!{RCOLS[j]}8"
            asm[f"{ac}{arow}"].font = Font(color=GREEN)
    say(f"거시 RAW 분리 {len(MACRO)}시트 — Assumption 하드값 이사 + 참조 전환")


def build_format(wb) -> None:
    """서식 규약을 문서가 아니라 워크북에 — 파일만 받아도 규약을 알 수 있게."""
    ws = wb.create_sheet("Format")
    ws.sheet_properties.tabColor = "94A3B8"
    title_of(ws, "서식 규약 (Val-Studio)")
    ws["B4"], ws["C4"], ws["D4"] = "구분", "예시", "의미"
    for c in ("B4", "C4", "D4"):
        ws[c].font = Font(bold=True, color="FFFFFFFF")
        ws[c].fill = PatternFill("solid", fgColor=HEAD)
    rows = [("파랑", BLUE, "직접 입력(hard number). 이 색이 아닌 셀은 건드리지 않는다."),
            ("검정", BLACK, "해당 시트 내 계산 수식."),
            ("초록", GREEN, "타 시트 참조. 참조 방향은 단방향(뒤→앞), 순환 금지.")]
    for i, (name, color, desc) in enumerate(rows, start=5):
        ws.cell(row=i, column=2, value=name)
        c = ws.cell(row=i, column=3, value="1,234")
        c.font = Font(color=color, bold=True)
        ws.cell(row=i, column=4, value=desc)
    ws["B10"] = "핵심 가정은 노랑 채움(fill)으로 추가 표시."
    ws["B10"].fill = PatternFill("solid", fgColor=YEL)
    ws["B12"] = "hard number 승격 규칙: 상류 시트가 생기면 하류의 파랑 입력셀을 초록 참조로 교체하고,"
    ws["B13"] = "교체 전후 주당가치가 불변인지(tie-out) 확인한다."
    for w, col in ((10, "B"), (12, "C"), (72, "D")):
        ws.column_dimensions[col].width = w


def build_separators(wb) -> None:
    for name, head, sub in (
        ("RESEARCH >>", "── 조사(RESEARCH) 구역 ──",
         "여기부터는 판단 근거·원자료. 모델 시트는 이 구역을 참조만 한다."),
        ("RAW >>", "── 원자료(RAW) 구역 ──",
         "r* 시트는 '붙여넣은 원자료' 전용. 가공·추정은 상류 시트에서 한다. "
         "hard number 는 RAW 에 1곳만."),
    ):
        ws = wb.create_sheet(name)
        ws.sheet_properties.tabColor = "94A3B8"
        ws["B2"] = head
        ws["B2"].font = Font(bold=True, size=12)
        ws["B3"] = sub


# ═══════════════════════════════════════════════════════════════════════════
# 페이즈 ⑤ 배선 — 세부 → 요약 → DCF 스파인
# ═══════════════════════════════════════════════════════════════════════════
def phase_wire(wb) -> int:
    """원본 빈 템플릿은 **하향 참조만** 배선돼 있다(WC/EBIT/FA 가 DCF 헤더·드라이버를 가져감).
    추정 결과가 DCF 스파인으로 올라가는 **상향 배선이 통째로 비어** 있어 그것을 채운다.
    추정기간 M:Q 에만 건다 — H:L(실적)·R(터미널) 원본 수식을 덮으면 모델이 깨진다."""
    ebit, fa, wc, dcf = wb["EBIT"], wb["FA"], wb["WC"], wb["DCF"]
    n = 0

    # EBIT 매출: 세부(판매량×단가) → 사업별 요약 → 합계
    for detail, qty, price in ((39, 40, 42), (44, 45, 47), (49, 50, 52), (54, 55, 57)):
        for c in FCOLS:
            n += put(ebit, f"{c}{detail}", f"={c}{qty}*{c}{price}")
    for c in FCOLS:
        n += put(ebit, f"{c}59", f"=SUM({c}39,{c}44,{c}49,{c}54)")
    for summ, detail in ((26, 39), (28, 44), (30, 49), (32, 54)):
        for c in FCOLS:
            n += put(ebit, f"{c}{summ}", f"={c}{detail}", GREEN)
    for c in FCOLS:
        n += put(ebit, f"{c}34", f"=SUM({c}26,{c}28,{c}30,{c}32)")
        # 매출원가(성격별 CASE2) · 판관비 합계
        n += put(ebit, f"{c}83", f"=SUM({c}78:{c}82)")
        n += put(ebit, f"{c}111", f"=SUM({c}88,{c}91,{c}95,{c}105,{c}108)")
        n += put(ebit, f"{c}123", f"=SUM({c}115,{c}117,{c}119,{c}121)")
        n += put(ebit, f"{c}148", f"=SUM({c}128,{c}132,{c}142,{c}145)")
        # EBIT 요약 블록
        n += put(ebit, f"{c}13", f"={c}34", GREEN)
        n += put(ebit, f"{c}15", f"={c}83", GREEN)
        n += put(ebit, f"{c}17", f"={c}13-{c}15")
        n += put(ebit, f"{c}19", f"={c}123", GREEN)
        n += put(ebit, f"{c}21", f"={c}17-{c}19")
        # FA 요약
        n += put(fa, f"{c}7", f"={c}8+{c}9")
        n += put(fa, f"{c}10", f"={c}11+{c}12")
        # DCF 스파인 ← 상류 (부호 규약: (+)D&A, (-)CAPEX, ΔNWC 현금조정)
        n += put(dcf, f"{c}7", f"=EBIT!{c}13", GREEN)
        n += put(dcf, f"{c}9", f"=EBIT!{c}15", GREEN)
        n += put(dcf, f"{c}13", f"=EBIT!{c}19", GREEN)
        n += put(dcf, f"{c}22", f"=FA!{c}7", GREEN)
        n += put(dcf, f"{c}23", f"=-FA!{c}10", GREEN)
        n += put(dcf, f"{c}24", f"=-WC!{c}14", GREEN)
    say(f"상향 배선 {n}건 (EBIT/FA → DCF 스파인)")
    return n


def phase_wire_labels(wb) -> None:
    """Research 부문명 → EBIT 사업1~4 라벨. 빈 템플릿에서도 깨지지 않도록 IF 폴백."""
    ebit = wb["EBIT"]
    RES_ROWS = [78, 79, 80, 81]                  # Research 제품표 데이터 1~4행
    for i, (summ, detail) in enumerate([(26, 39), (28, 44), (30, 49), (32, 54)]):
        f = f'=IF(Research!C{RES_ROWS[i]}="","사업{i+1} 매출",Research!C{RES_ROWS[i]}&" 매출")'
        for cell in (f"C{summ}", f"C{detail}"):
            ebit[cell] = f
            ebit[cell].font = Font(color=GREEN)
    say("라벨 배선 8건 (Research 부문명 → EBIT 사업1~4)")


def phase_wacc(wb) -> None:
    """원본 WACC 시트는 라벨만 있고 계산이 없다 — 표준 CAPM 빌드업을 심고 DCF!H37 로 승격.

    D40(D/(D+E))을 입력이 아니라 D21(D/E)에서 **유도**한다 — 둘 다 입력받으면 서로
    어긋나는 순간 WACC 가 조용히 틀린다.
    """
    w, dcf = wb["WACC"], wb["DCF"]
    calcs = {
        23: "=D20*(1+(1-D22)*D21)",                 # Hamada 재레버
        27: "=D23",
        33: "=D25+D27*D26+D29+D30+D31",             # Ke = Rf + β·MRP + Size + CRP + CSR
        36: "=D22",
        38: "=D35*(1-D36)",                         # 세후 Kd
        40: "=IFERROR(D21/(1+D21),0)",
        41: "=1-D40",
        43: "=D33*D41+D38*D40",
    }
    for r in (20, 21, 22, 25, 26, 29, 30, 31, 35):  # 파랑 입력 9칸
        if w[f"D{r}"].value is None:
            w[f"D{r}"].font = Font(color=BLUE)
    for r, f in calcs.items():
        if put(w, f"D{r}", f, GREEN if r == 27 else BLACK, bold=(r == 43)):
            w[f"D{r}"].number_format = "0.00%"
    w["E43"] = "← DCF!H37 이 이 셀을 참조한다"
    w["E43"].font = Font(size=9, color=GREY)
    old = dcf["H37"].value
    dcf["H37"] = "=WACC!D43"                        # hard number 승격(SSOT 이동)
    dcf["H37"].font = Font(color=GREEN)
    say(f"WACC 빌드업 8수식 + DCF!H37 승격(기존 하드 {old})")


# ═══════════════════════════════════════════════════════════════════════════
# 페이즈 ⑥ 메타 — _DART_MAP · _VS_STATE
# ═══════════════════════════════════════════════════════════════════════════
# (대상시트, 섹션, 항목, DART 원천, 엔드포인트, 등급)
# A=구조화 응답으로 바로 채움 / B=문서 표 추출 필요 / C=서술형(AI 제안→평가인 확인)
DART_MAP = [
    ("Research", "2.회사개요", "설립일·대표이사·본사·홈페이지", "기업개황", "/api/dart/company", "A"),
    ("Research", "2.회사개요", "상장시장·상장일·결산월", "기업개황", "/api/dart/company", "A"),
    ("Research", "2.회사개요", "감사인·감사의견", "감사보고서", "/api/dart/audit-opinion", "A"),
    ("Research", "2.회사개요", "발행주식수·자기주식·유통주식", "주식총수 현황", "/api/dart/shares", "A"),
    ("Research", "2.회사개요", "주주 구성·최대주주 지분율", "최대주주 현황", "/api/dart/investments", "A"),
    ("Research", "2.회사개요", "주요 종속회사·지분율", "타법인 출자현황", "/api/dart/investments", "A"),
    ("Research", "2.회사개요", "배당성향", "배당에 관한 사항", "/api/dart/dividends", "A"),
    ("Research", "2.회사개요", "직원수", "임직원 현황", "/api/dart/employee", "A"),
    ("Research", "4.제품별매출", "사업부문별 매출·비중", "II.사업의 내용(매출 실적)",
     "/api/dart/document → 표추출", "B"),
    ("Research", "4.제품별매출", "종속회사별 제품 현황", "II.사업의 내용",
     "/api/dart/document → 표추출", "B"),
    ("Research", "9.고객·매출처", "주요 고객·매출 비중", "II.사업의 내용(주요 매출처)",
     "/api/dart/document → 표추출", "B"),
    ("Research", "5.주요제품설명", "제품명·용도·전방산업", "II.사업의 내용(제품 설명)",
     "/api/dart/document → 서술 추출", "C"),
    ("Research", "5.주요제품설명", "시장 점유율", "II.사업의 내용(시장 여건)",
     "/api/dart/document → 서술 추출", "C"),
    ("Research", "6.ValueChain", "전·후방산업, 진입장벽", "II.사업의 내용",
     "/api/dart/document → 서술 추출", "C"),
    ("rFS", "재무제표", "재무상태표·손익계산서·현금흐름표 원문", "재무제표(XBRL)",
     "/api/dart/financials", "A"),
    ("rFS", "재무제표", "세그먼트·주식수 프리필", "XBRL", "/api/brief/from_xbrl", "A"),
    ("Peer_S1_모집단", "모집단", "업종별 상장사 목록", "-", "/api/ksic/search", "A"),
    ("Comps", "유사회사", "peer 4-step 퍼널 결과", "-", "/api/peer/select", "A"),
    ("rTrading", "시장데이터", "주가·시총·발행주식수", "-",
     "/api/price/marketcap · /api/price/multiples", "A"),
    ("rGDP", "거시", "실질 GDP 성장률", "-", "/api/macro/series (ECOS)", "A"),
    ("rInflation", "거시", "소비자물가 상승률", "-", "/api/macro/series (ECOS)", "A"),
    ("rWage", "거시", "명목임금 상승률", "-", "/api/macro/series", "A"),
]


def build_dart_map(wb) -> None:
    """인제스트 코드가 셀 주소를 하드코딩하지 않도록 원천 계약을 워크북에 선언."""
    ws = wb.create_sheet("_DART_MAP")
    ws.sheet_properties.tabColor = "94A3B8"
    title_of(ws, "자동수집 매핑 선언 (_DART_MAP)",
             "어느 시트·항목이 어느 원천에서 오는지의 계약. 현재는 선언만 — 호출 배관은 별도.")
    ws["B4"] = ("등급  A=구조화 응답으로 바로 채움 · B=문서 표 추출 필요 · "
                "C=서술형(AI 제안 → 평가인 확인)")
    ws["B4"].font = Font(size=9, color=GREY)
    heads = ["대상 시트", "섹션", "항목", "DART 원천", "엔드포인트", "등급", "상태"]
    for j, (h, w) in enumerate(zip(heads, [18, 16, 32, 28, 40, 6, 10])):
        c = ws.cell(row=6, column=2 + j, value=h)
        c.font = Font(bold=True, size=9, color="FFFFFFFF")
        c.fill = PatternFill("solid", fgColor=HEAD)
        c.alignment = Alignment(horizontal="center")
        c.border = BOX
        ws.column_dimensions[chr(ord("B") + j)].width = w
    for i, row in enumerate(DART_MAP, start=7):
        for j, v in enumerate([*row, "미배관"]):
            c = ws.cell(row=i, column=2 + j, value=v)
            c.border = BOX
            c.font = Font(size=9)
            if j == 5:
                c.alignment = Alignment(horizontal="center")
                c.font = Font(size=9, bold=True,
                              color={"A": GREEN, "B": "FFB45309", "C": "FFB91C1C"}[v])
    r = 7 + len(DART_MAP) + 1
    cnt = {g: sum(1 for *_, gg in DART_MAP if gg == g) for g in "ABC"}
    ws.cell(row=r, column=2,
            value=f"합계 {len(DART_MAP)}건 — A {cnt['A']} · B {cnt['B']} · C {cnt['C']}").font = Font(bold=True)
    ws.cell(row=r + 1, column=2,
            value="⚠ 등급 C(서술형)는 평가인 확인 전까지 하류 가정으로 흘려보내지 않는다.").font = Font(size=9, color=GREY)
    say(f"_DART_MAP {len(DART_MAP)}건 (A {cnt['A']}/B {cnt['B']}/C {cnt['C']})")


def build_vs_state(wb, n_years: int = 5) -> None:
    """워크북=상태 규약. `layout` 키가 되읽기 라우팅(fullmodel_layout)의 1순위 판별 근거."""
    ws = wb.create_sheet("_VS_STATE")
    kv = [("skill_version", "1.0"), ("mode", "A"), ("stage", "W0"),
          ("layout", LAYOUT_ID), ("last_gate_passed", "W0:full-template"),
          ("n_years", n_years)]
    for i, (k, v) in enumerate(kv, start=1):
        ws.cell(row=i, column=1, value=k)
        ws.cell(row=i, column=2, value=v)
    hdr_row = len(kv) + 2
    ws.cell(row=hdr_row, column=1, value="── 가정 대장(provenance) ──")
    for col, label in zip("ABCDE", ["가정명", "값", "출처유형", "근거", "승인상태"]):
        ws[f"{col}{hdr_row + 1}"] = label
        ws[f"{col}{hdr_row + 1}"].font = Font(bold=True)
    ws.sheet_state = "hidden"
    say(f"_VS_STATE (layout={LAYOUT_ID}, 숨김)")


# ═══════════════════════════════════════════════════════════════════════════
# 검증 — 배선·순환·판별
# ═══════════════════════════════════════════════════════════════════════════
TOKEN = re.compile(r"(?:('[^']+'|[A-Za-z_가-힣][\w가-힣 .>]*)!)?\$?([A-Z]{1,3})\$?(\d{1,5})")


def verify(path: Path) -> bool:
    wb = openpyxl.load_workbook(path, data_only=False)
    ok = True
    print("\n[검증]")

    checks = [("DCF", "M7", "=EBIT!M13"), ("DCF", "M22", "=FA!M7"), ("DCF", "M24", "=-WC!M14"),
              ("DCF", "H37", "=WACC!D43"), ("WACC", "D43", "=D33*D41+D38*D40"),
              ("Assumption", "E9", "=rGDP!C8"), ("EBIT", "M13", "=M34")]
    for s, ref, want in checks:
        got = wb[s][ref].value
        good = got == want
        ok &= good
        print(f"  {s}!{ref:<5} {str(got):<26} {'OK' if good else f'!= {want}'}")

    # 셀 단위 순환(Excel 은 시트가 아니라 셀 단위로 판정 — 시트 간 상호참조 자체는 무해)
    sheets = set(wb.sheetnames)
    graph = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    deps = {((sh or "").strip("'") or ws.title, f"{col}{r}")
                            for sh, col, r in TOKEN.findall(c.value[1:])}
                    graph[(ws.title, c.coordinate)] = {d for d in deps if d[0] in sheets}
    sys.setrecursionlimit(20000)
    color, cycles = {}, []

    def dfs(u, stack):
        color[u] = 1
        stack.append(u)
        for v in graph.get(u, ()):
            if v not in graph:
                continue
            if color.get(v, 0) == 1:
                cycles.append(stack[stack.index(v):] + [v])
            elif color.get(v, 0) == 0:
                dfs(v, stack)
        stack.pop()
        color[u] = 2

    for node in list(graph):
        if color.get(node, 0) == 0:
            dfs(node, [])
    ok &= not cycles
    print(f"  수식 {len(graph):,}개 · 순환 {len(cycles)}건")
    for cyc in cycles[:3]:
        print("   " + " → ".join(f"{s}!{r}" for s, r in cyc))

    # 레이아웃 판별(백엔드 어댑터가 이 파일을 인식하는가)
    try:
        from excel.fullmodel_layout import detect_fullmodel
        from excel.xlsx_reader import read_workbook
        det = detect_fullmodel(read_workbook(str(path)))
        ok &= det
        print(f"  detect_fullmodel = {det}")
    except ImportError:
        print("  detect_fullmodel = (백엔드 미로드 — 건너뜀)")

    leftover = sum(len(BRAND.findall(str(c.value)))
                   for ws in wb.worksheets for row in ws.iter_rows()
                   for c in row if isinstance(c.value, str))
    ok &= leftover == 0
    print(f"  브랜딩 잔존 {leftover}건")
    print(f"\n판정: {'PASS' if ok else 'FAIL'}")
    return ok


# ═══════════════════════════════════════════════════════════════════════════
def build(out: Path) -> None:
    missing = [p for p in (SRC_BASE, SRC_PEER, SRC_COMPS) if not p.exists()]
    if missing:
        raise SystemExit("원본 서식 없음 — 조용히 반쪽 템플릿을 만들지 않는다:\n  "
                         + "\n  ".join(str(p) for p in missing))

    print(f"[build] {out.name}")
    wb = openpyxl.load_workbook(SRC_BASE, data_only=False)
    say(f"베이스 로드 {len(wb.sheetnames)}시트")

    phase_clean(wb)
    phase_peer(wb)
    phase_comps(wb)
    build_research(wb)
    build_rfs(wb)
    build_macro_raw(wb)
    build_format(wb)
    build_separators(wb)
    phase_wire(wb)
    phase_wire_labels(wb)
    phase_wacc(wb)
    build_dart_map(wb)
    build_vs_state(wb)

    order = [s for s in TAB_ORDER if s in wb.sheetnames]
    order += [s for s in wb.sheetnames if s not in order]
    wb._sheets = [wb[s] for s in order]
    wb.save(out)
    say(f"저장 {out} ({out.stat().st_size/1024/1024:.2f} MB · {len(wb.sheetnames)}시트)")


def main() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except (AttributeError, ValueError):
        pass
    args = sys.argv[1:]
    out = OUT_DEFAULT
    if "--out" in args:
        out = Path(args[args.index("--out") + 1])
    if "--verify-only" not in args:
        build(out)
    if not verify(out):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
